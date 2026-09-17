import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).parent
SOURCE = ROOT.parent / "outputs/01a05ac3-4162-75b3-9fea-ebcd8b05da27/Susan基础句型与SSS儿歌打卡拆解版_小小优趣收录_33周打卡_直接跳转歌曲Sheet.xlsx"


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def sheet_from_location(location):
    if not location:
        return ""
    match = re.match(r"^'(.*)'![A-Z]+\d+$", location)
    return match.group(1).replace("''", "'") if match else ""


def formula_url(value):
    if not isinstance(value, str):
        return ""
    match = re.search(r'HYPERLINK\("([^"]+)"', value, re.I)
    return match.group(1) if match else ""


wb = load_workbook(SOURCE, read_only=False, data_only=False)
cached = load_workbook(SOURCE, read_only=True, data_only=True)

tracker = wb["33周儿歌打卡"]
tracker_cached = cached["33周儿歌打卡"]
weeks = []
song_sheet_by_title = {}
for row_num in range(6, 39):
    songs = []
    for title_col, status_col in ((3, 4), (5, 6), (7, 8)):
        cell = tracker.cell(row_num, title_col)
        title = clean(cell.value)
        sheet_name = sheet_from_location(cell.hyperlink.location if cell.hyperlink else "")
        status = clean(tracker_cached.cell(row_num, status_col).value) or "待打卡"
        songs.append({"title": title, "sheet": sheet_name, "status": status})
        if title and sheet_name:
            song_sheet_by_title.setdefault(title.casefold(), sheet_name)
    weeks.append({
        "week": row_num - 5,
        "label": clean(tracker.cell(row_num, 1).value),
        "theme": clean(tracker.cell(row_num, 2).value),
        "songs": songs,
    })

details = wb["歌曲明细"]
details_cached = cached["歌曲明细"]
detail_rows = []
for row_num in range(7, details.max_row + 1):
    title_cell = details.cell(row_num, 4)
    title = clean(title_cell.value)
    if not title:
        continue
    sheet_name = sheet_from_location(title_cell.hyperlink.location if title_cell.hyperlink else "")
    if not sheet_name:
        sheet_name = song_sheet_by_title.get(title.casefold(), "")
    detail_rows.append({
        "type": clean(details.cell(row_num, 1).value),
        "pattern": clean(details.cell(row_num, 2).value),
        "example": clean(details.cell(row_num, 3).value),
        "songTitle": title,
        "sheet": sheet_name,
        "category": clean(details.cell(row_num, 5).value),
        "videoStatus": clean(details.cell(row_num, 6).value),
        "videoUrl": clean(details.cell(row_num, 7).value),
        "patternStatus": clean(details_cached.cell(row_num, 8).value) or "待打卡",
        "songStatus": clean(details_cached.cell(row_num, 9).value) or "待打卡",
        "ukidsStatus": clean(details.cell(row_num, 10).value),
        "ukidsUrl": formula_url(details.cell(row_num, 11).value),
    })

detail_by_sheet = {}
for item in detail_rows:
    if item["sheet"]:
        detail_by_sheet.setdefault(item["sheet"], []).append(item)


def find_section_rows(ws):
    result = {}
    for row_num in range(1, ws.max_row + 1):
        value = clean(ws.cell(row_num, 1).value)
        if value.startswith("1. 完整歌词"):
            result["lyrics"] = row_num
        elif value.startswith("2. 核心词汇"):
            result["words"] = row_num
        elif value.startswith("3. TPR"):
            result["tpr"] = row_num
        elif value.startswith("4. 原表匹配句型"):
            result["patterns"] = row_num
    return result


songs = []
for sheet_name in wb.sheetnames:
    if not re.search(r" \d{3}$", sheet_name):
        continue
    ws = wb[sheet_name]
    ws_cached = cached[sheet_name]
    sections = find_section_rows(ws)
    if not {"lyrics", "words", "tpr", "patterns"}.issubset(sections):
        continue

    lyrics = []
    for row_num in range(sections["lyrics"] + 1, sections["words"]):
        line = clean(ws.cell(row_num, 1).value)
        if line:
            lyrics.append(line)

    words = []
    for row_num in range(sections["words"] + 2, sections["tpr"]):
        english = clean(ws.cell(row_num, 1).value)
        if english:
            words.append({
                "english": english,
                "ipa": clean(ws.cell(row_num, 2).value),
                "meaning": clean(ws.cell(row_num, 3).value),
            })

    tpr = []
    for row_num in range(sections["tpr"] + 2, sections["patterns"]):
        step = clean(ws.cell(row_num, 1).value)
        lyric = clean(ws.cell(row_num, 2).value)
        translation = clean(ws.cell(row_num, 5).value)
        if step or lyric or translation:
            tpr.append({"step": step, "lyric": lyric, "translation": translation})

    matches = detail_by_sheet.get(sheet_name, [])
    meta = clean(ws.cell(2, 1).value)
    title = clean(ws.cell(1, 1).value) or re.sub(r" \d{3}$", "", sheet_name)
    first_match = matches[0] if matches else {}
    video_url = first_match.get("videoUrl", "")
    ukids_url = first_match.get("ukidsUrl", "")
    status = first_match.get("songStatus", "待打卡")
    song_id_match = re.search(r"(\d{3})$", sheet_name)
    songs.append({
        "id": song_id_match.group(1) if song_id_match else sheet_name,
        "sheet": sheet_name,
        "title": title,
        "meta": meta,
        "category": first_match.get("category", meta.split("|")[0].strip()),
        "defaultStatus": status,
        "lyrics": lyrics,
        "words": words,
        "tpr": tpr,
        "patterns": [{
            "type": row["type"],
            "pattern": row["pattern"],
            "example": row["example"],
            "status": row["patternStatus"],
        } for row in matches],
        "videoUrl": video_url,
        "ukidsStatus": first_match.get("ukidsStatus", ""),
        "ukidsUrl": ukids_url,
    })

week_lookup = {}
for week in weeks:
    for item in week["songs"]:
        if item["sheet"]:
            week_lookup[item["sheet"]] = {"week": week["week"], "theme": week["theme"]}
for song in songs:
    song.update(week_lookup.get(song["sheet"], {}))

payload = {"weeks": weeks, "songs": songs, "detailRows": detail_rows}
(ROOT / "data.js").write_text(
    "window.SSS_WEB_DATA = " + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n",
    encoding="utf-8",
)

print(json.dumps({
    "weeks": len(weeks),
    "songs": len(songs),
    "detailRows": len(detail_rows),
    "lyrics": sum(len(song["lyrics"]) for song in songs),
    "words": sum(len(song["words"]) for song in songs),
    "tpr": sum(len(song["tpr"]) for song in songs),
}, ensure_ascii=False))

wb.close()
cached.close()
