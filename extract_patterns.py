import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/sss_pattern_convert/基本句型.xlsx")
OUTPUT = Path(__file__).with_name("pattern-data.js")
SHEETS = {
    "名词基础句型": ("名词卡", "A", 51),
    "动词基础句型": ("动词卡", "B", 68),
    "形容词基础句型": ("形容词卡", "C", 100),
}


workbook = load_workbook(SOURCE, read_only=True, data_only=True)
items = []
for sheet_name, (type_name, prefix, maximum) in SHEETS.items():
    worksheet = workbook[sheet_name]
    for row in worksheet.iter_rows(min_row=4, values_only=True):
        grammar, code, pattern, example, notes = row[1:6]
        match = re.fullmatch(rf"{prefix}0*(\d+)", str(code or "").strip(), re.IGNORECASE)
        if not match or not pattern:
            continue
        number = int(match.group(1))
        if number > maximum:
            continue
        items.append({
            "type": type_name,
            "number": number,
            "code": f"{prefix}{number:02d}" if prefix != "A" else f"{prefix}{number}",
            "grammar": str(grammar or "").strip(),
            "pattern": str(pattern).strip(),
            "example": str(example or "").strip(),
            "notes": str(notes or "").strip(),
        })

items.sort(key=lambda item: (list(SHEETS).index(next(name for name, value in SHEETS.items() if value[0] == item["type"])), item["number"]))
OUTPUT.write_text("window.SSS_PATTERN_LIBRARY = " + json.dumps(items, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
print(json.dumps({"output": str(OUTPUT), "count": len(items), "by_type": {name: sum(item["type"] == name for item in items) for name, _, _ in SHEETS.values()}}, ensure_ascii=False))
