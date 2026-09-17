import json
from pathlib import Path
from openpyxl import load_workbook

SOURCE = Path("/Users/jingjie/Desktop/未命名文件夹 2/副本名词引导输出句型(1)(1).xlsx")
OUTPUT = Path(__file__).with_name("noun-detail-data.js")
EXTRAS = [
    "It's a tiger.", "This is my cup.", "That is an elephant.", "A tomato is red.", "The sea is blue.",
    "It's a yellow duck.", "The whale is big.", "An ant is small.", "It isn't my bag.", "I can see a kite.",
    "I can see three ducks.", "It's my teddy bear.", "She is my teacher.", "This pen isn't mine.", "That's Lucy's doll.",
    "There is a cat on the sofa.", "There is a whale in the sea.", "There is a bee on the grass.", "I can't see the moon.", "I have two books.",
    "I don't have a bike.", "My dad has a car.", "My dad doesn't have a red coat.", "I like bananas.", "I don't like onions.",
    "The rabbit likes carrots.", "The cat doesn't like water.", "I love story time.", "I hate loud noises.", "Grandpa loves his garden.",
    "Grandma hates cold weather.", "I want some water.", "I don't want soup.", "The dog wants a bone.", "Mr. Bull doesn't want the red hat.",
    "Can I have some water, please?", "I want to eat an apple.", "I am wearing blue socks.", "She is wearing a red dress.", "A rabbit doesn't eat meat.",
    "I can drink orange juice.", "Look at the rainbow.", "Listen to the rain.", "I can hear Mom singing.", "The soup is hot.",
    "The ice is cold.", "The cup is on the table.", "The peas are in the freezer.", "Your shoes are under the bed.", "The rainbow is beautiful.", "My cookie is gone.",
]

worksheet = load_workbook(SOURCE, read_only=True, data_only=True).active
headers = [str(value or "").strip() for value in next(worksheet.iter_rows(min_row=2, max_row=2, values_only=True))]
items = []
for row in worksheet.iter_rows(min_row=3, values_only=True):
    if not row[0] or not row[1]:
        continue
    number = int(row[0])
    forms = [{"label": headers[index], "value": str(row[index]).strip()} for index in range(2, 10) if row[index]]
    items.append({"number": number, "code": f"A{number}", "pattern": str(row[1]).strip(), "forms": forms, "extraCases": [EXTRAS[number - 1]]})

OUTPUT.write_text("window.SSS_NOUN_DETAILS = " + json.dumps(items, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
print(json.dumps({"count": len(items), "output": str(OUTPUT)}, ensure_ascii=False))
